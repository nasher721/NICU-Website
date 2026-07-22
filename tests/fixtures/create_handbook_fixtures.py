#!/usr/bin/env python3
"""Create small clean DOCX sources for extractor regeneration tests."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def create_handbook(path: Path, campus: str):
    document = Document()
    document.add_paragraph(f"{campus} faculty orientation fixture")
    document.add_heading("Daily Workflow", level=1)
    document.add_paragraph("Rounds and service expectations are source-controlled.")
    document.add_heading("Documentation Quality", level=2)
    document.add_paragraph("Use the reviewed documentation workflow.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Task"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Review"
    table.cell(1, 1).text = campus
    document.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    create_handbook(args.output / "main-campus.docx", "Main Campus")
    create_handbook(args.output / "akron.docx", "Akron General")


if __name__ == "__main__":
    main()
